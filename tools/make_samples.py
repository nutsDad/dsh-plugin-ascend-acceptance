"""生成用于插件自测的 DOCX / XLSX 示例（仅开发时运行）。"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"
SAMPLES.mkdir(exist_ok=True)

FONT_CN = "Microsoft YaHei"
ACCENT = "C2410C"


def set_run_font(run, size=10.5, bold=False, color=None):
    run.font.name = FONT_CN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_table(document, rows, widths=None):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        for cell_index, value in enumerate(row):
            cell = table.cell(row_index, cell_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(str(value))
            set_run_font(run, 9.5, bold=row_index == 0, color="FFFFFF" if row_index == 0 else None)
            if row_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = width
    return table


def build_standard():
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Pt(46)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(10.5)

    title = document.add_heading("昇腾智算集群交付标准 V2.1（示例）", level=0)
    for run in title.runs:
        set_run_font(run, 20, bold=True, color=ACCENT)
    paragraph = document.add_paragraph("适用于昇腾智算集群到货验收；示例数据仅用于插件自测。")
    for run in paragraph.runs:
        set_run_font(run, 9.5, color="666666")

    document.add_heading("1 昇腾运行环境", level=1)
    add_table(document, [
        ["序号", "指标", "交付要求", "验证方法"],
        ["1", "NPU 型号", "Atlas 800T A2 或同等级昇腾训练服务器", "到货清单 + npu-smi info"],
        ["2", "单节点 NPU 数量", "不少于 8 张", "npu-smi info -l"],
        ["3", "NPU 健康状态", "全部正常", "npu-smi info -t health"],
        ["4", "NPU 温度", "不高于 85 ℃", "npu-smi info -t temp"],
        ["5", "NPU 功耗", "不高于 400 W", "npu-smi info -t power"],
        ["6", "NPU 显存", "不小于 64 GB", "npu-smi info -t memory"],
        ["7", "主机内存", "不小于 1024 GB", "free -h"],
        ["8", "ECC 错误计数", "等于 0", "npu-smi info -t ecc"],
    ], widths=[Pt(38), Pt(110), Pt(220), Pt(150)])
    for text in [
        "CANN 版本应不低于 8.0.RC1。",
        "NPU 驱动版本应不低于 23.0.3。",
        "NPU 固件版本应不低于 7.1.0.5.220。",
        "操作系统应为 openEuler 22.03 LTS 或以上。",
        "Python 版本应不低于 3.10。",
    ]:
        paragraph = document.add_paragraph(text)
        for run in paragraph.runs:
            set_run_font(run)

    document.add_heading("2 机房基础设施", level=1)
    add_table(document, [
        ["序号", "指标", "交付要求", "验证方法"],
        ["1", "机房温度", "冷通道温度保持在 18 ℃～27 ℃", "动环监控 24h 曲线"],
        ["2", "相对湿度", "保持在 30%～70% RH", "动环监控 24h 曲线"],
        ["3", "单机柜功率", "不低于 8 kW", "配电核算 + 实测电流"],
        ["4", "机柜高度", "不小于 42U", "机柜铭牌"],
        ["5", "楼板承重", "不低于 1000 kg/m²", "结构复核报告"],
        ["6", "接地电阻", "不大于 4 Ω", "接地电阻测试报告"],
        ["7", "UPS 后备时间", "不少于 30 分钟", "UPS 放电测试报告"],
    ], widths=[Pt(38), Pt(110), Pt(220), Pt(150)])
    for text in [
        "机房应配置双路市电和 A/B 路 PDU，单路故障时不影响业务运行。",
        "机房应配置气体灭火系统，并完成消防联动测试。",
        "机房应配置漏水检测，接入动环监控并支持短信告警。",
        "冷通道应完成封闭，避免冷热气流短路。",
    ]:
        paragraph = document.add_paragraph(text)
        for run in paragraph.runs:
            set_run_font(run)

    document.add_heading("3 交付文档与运维", level=1)
    for text in [
        "应提供以下交付文档：到货验收单、装箱单、出厂检测报告、产品合格证、保修卡、资产清单、网络拓扑图、操作手册。",
        "应提供集群性能测试报告，包含算力、网络和存储测试的原始数据。",
        "质保期不少于 3 年。",
        "故障响应时间不超过 4 小时。",
    ]:
        paragraph = document.add_paragraph(text)
        for run in paragraph.runs:
            set_run_font(run)

    output = SAMPLES / "昇腾交付标准-示例.docx"
    document.save(output)
    return output


def build_delivery():
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(name=FONT_CN, size=10, bold=True, color="FFFFFF")
    body_font = Font(name=FONT_CN, size=10)

    def sheet_rows(sheet, headers, rows):
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column, width in zip("ABCD", [24, 30, 42, 20]):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"

    sheet = workbook.active
    sheet.title = "运行环境"
    sheet_rows(sheet, ["指标", "实测值", "证据", "备注"], [
        ["NPU 型号", "Atlas 800T A2", "到货装箱单、设备铭牌", ""],
        ["单节点 NPU 数量", "8 张", "npu-smi info -l", ""],
        ["NPU 健康状态", "全部正常", "npu-smi info -t health", ""],
        ["NPU 温度", "78 ℃", "满载 30 分钟", ""],
        ["NPU 功耗", "365 W", "npu-smi info -t power", ""],
        ["NPU 显存", "64 GB", "npu-smi info -t memory", ""],
        ["主机内存", "1024 GB", "free -h", ""],
        ["ECC 错误计数", "0", "npu-smi info -t ecc", ""],
        ["CANN 版本", "8.0.RC2", "version.cfg", ""],
        ["NPU 驱动版本", "23.0.3", "npu-smi info", ""],
        ["NPU 固件版本", "7.1.0.5.220", "npu-smi info -t sysinfo", ""],
        ["操作系统版本", "openEuler 22.03 LTS", "/etc/os-release", ""],
        ["Python 版本", "3.10.12", "python3 --version", ""],
    ])

    sheet = workbook.create_sheet("机房")
    sheet_rows(sheet, ["指标", "实测值", "证据", "备注"], [
        ["机房温度", "23.5 ℃", "动环 24h 曲线，最高 25.8 ℃", ""],
        ["相对湿度", "45% RH", "动环 24h 曲线", ""],
        ["单机柜功率", "6 kW", "配电核算 + 实测", "低于标准"],
        ["机柜高度", "47U", "机柜铭牌", ""],
        ["楼板承重", "800 kg/m²", "结构复核报告", "低于标准"],
        ["接地电阻", "3.2 Ω", "接地电阻测试报告", ""],
        ["UPS 后备时间", "45 分钟", "放电测试报告", ""],
        ["PUE", "1.62", "连续 30 天计量", "高于标准"],
        ["供电方式", "双路市电 + A/B 路 PDU", "已完成单路切换测试", ""],
        ["气体灭火", "已验收", "消防联动测试记录", ""],
        ["漏水检测", "已接入动环", "模拟告警正常", ""],
        ["冷通道", "已完成封闭", "现场照片见附件 9", ""],
    ])

    sheet = workbook.create_sheet("运维保障")
    sheet_rows(sheet, ["指标", "实测值", "证据", "备注"], [
        ["质保期", "3 年", "原厂质保函", ""],
        ["故障响应时间", "8 小时", "SLA 服务承诺函", "高于标准"],
        ["备品备件", "已在项目所在地设置备件库", "备件清单见附件 12", ""],
        ["定期巡检", "每季度一次，最近一期报告已提供", "巡检报告", ""],
    ])

    sheet = workbook.create_sheet("交付文档")
    sheet.append(["文档名称", "状态", "说明"])
    for row in [
        ["到货验收单", "已提供", "双方签字"],
        ["装箱单", "已提供", ""],
        ["出厂检测报告", "已提供", ""],
        ["产品合格证", "已提供", ""],
        ["保修卡", "已提供", ""],
        ["资产清单", "已提供", ""],
        ["网络拓扑图", "已提供", ""],
        ["操作手册", "待补充", "原厂中文版预计 10 月 25 日提供"],
        ["集群性能测试报告", "已提供", "含算力、网络测试数据"],
        ["培训记录", "部分提供", "签到表与课件已提供，考核结果待补"],
    ]:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 48

    output = SAMPLES / "验收交付件-示例.xlsx"
    workbook.save(output)
    return output


if __name__ == "__main__":
    print(build_standard())
    print(build_delivery())
